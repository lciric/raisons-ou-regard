import sys
from openpyxl import load_workbook
path = sys.argv[1]
wb = load_workbook(path)
for ws in wb.worksheets:
    print(f"=== sheet {ws.title!r} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} merged={list(ws.merged_cells.ranges)[:20]}")
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            fill = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
            color = c.font.color.rgb if c.font and c.font.color is not None else None
            v = str(c.value).replace("\n", "\\n")
            print(f"{c.coordinate:6} fill={fill} font={color} bold={c.font.bold} fmt={c.number_format!r} | {v[:160]}")
    dv = ws.data_validations.dataValidation if ws.data_validations else []
    for d in dv:
        print("  validation", d.type, d.formula1, d.sqref)
