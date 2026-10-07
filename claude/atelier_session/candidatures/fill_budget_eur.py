"""The merged EA Funds budget (sycophancy papers + "Reasons or being watched?"), in EUR, on EA Funds' own template:
three scenario tabs, only the template's input cells written (A-D from row 6, B1 rate, B2 contingency)."""
import sys
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

src, out = sys.argv[1], sys.argv[2]
wb = load_workbook(src)
tpl = wb["Template"]
RATE = "=1/1.1269"   # EUR per USD: ECB reference rate of 6 October 2026, 1 EUR = 1.1269 USD
H100 = "H100 SXM on vast.ai at 3.50 USD/hour (2.86-4.04 USD/h seen on 6 Oct 2026), converted at B1"

def stipend(months):
    return ("Personnel", f"Stipend for the applicant, gross: {months} months x 5,000 EUR",
            "Includes French income tax and social contributions, as the form requires. Proposed level, to be confirmed by the applicant", f"={months}*5000")

def gpu(hours, what, note):
    return ("Other", f"GPU compute, {what}: {hours:,} H100-hours", f"{note}. {H100}", f"={hours}*3.5*$B$1")

SYC_PAPERS = lambda hours, months: gpu(hours, "sycophancy papers (behaviour, measurement)",
    f"About 240 H100-hours a month, the rate of the earlier draft's compute line, over {months} months; 70B checks through NDIF at no cost")
RR_API = ("Other", "Claude API, 'Reasons or being watched?': training data of the six arms, cue sets, held-out scenarios, judges on subsamples",
          "Programme estimate in USD: 500-1,400 for the training data (upper end) and about 1,000 for the rest; lower with batch pricing, or by about 1,000 if Anthropic's External Researcher Access Program grants credits", "=2400*$B$1")
RESERVE = gpu(720, "reserve for the larger gaze test",
              "The gaze test grows from 150-280 to 600-1,000 GPU-h at 400 scenarios x 10 generations; spent only if the pilot's power simulation calls for it")

SCENARIOS = {
    "Minimum": [
        stipend(4),
        SYC_PAPERS(480, 2),
        gpu(480, "'Reasons or being watched?' minimal experiment", "Upper end of the programme's estimate for the minimal decisive experiment (230-480 GPU-h)"),
        ("Other", "LLM-judge API, sycophancy: judging the behavioural campaigns", "From the earlier draft's minimum scenario", 1200),
        RR_API,
        ("Operations", "Software and storage for sealed artefacts", "From the earlier draft", 200),
    ],
    "Mainline": [
        stipend(6),
        SYC_PAPERS(720, 3),
        gpu(560, "'Reasons or being watched?' first paper", "Upper end of the programme's estimate for the minimal experiment and the localization (270-560 GPU-h)"),
        RESERVE,
        ("Other", "LLM-judge API, sycophancy: judging and the validity battery", "From the earlier draft's mainline scenario", 2000),
        RR_API,
        ("Travel", "Conference travel to present the behavioural paper", "From the earlier draft", 2500),
        ("Operations", "Software and storage for sealed artefacts", "From the earlier draft", 300),
    ],
    "Maximum": [
        stipend(6),
        SYC_PAPERS(720, 3),
        gpu(560, "'Reasons or being watched?' first paper", "Upper end of the programme's estimate for the minimal experiment and the localization (270-560 GPU-h)"),
        RESERVE,
        ("Other", "LLM-judge API, sycophancy: judging and the validity battery", "From the earlier draft's mainline scenario", 2000),
        RR_API,
        ("Other", "70B runs on rented GPUs, only if NDIF cannot host the model", "From the earlier draft", 3000),
        ("Other", "Human-rated validation subset for the judge study", "From the earlier draft", 1500),
        ("Travel", "Conference travel: the behavioural paper and a second venue or research visit", "From the earlier draft", 4500),
        ("Operations", "Software and storage for sealed artefacts", "From the earlier draft", 300),
    ],
}

for name, lines in SCENARIOS.items():
    ws = wb.copy_worksheet(tpl)
    ws.title = name
    for col, dim in tpl.column_dimensions.items():
        ws.column_dimensions[col].width = dim.width
    dv = DataValidation(type="list", formula1='"Personnel,Materials,Operations,Travel,Other"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A6:A41")
    ws["B1"] = RATE
    ws["B2"] = 0.10      # contingency: the application form suggests about 10 %, the template 5 %
    for i, (kind, detail, note, amount) in enumerate(lines):
        r = 6 + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = kind, detail, note, amount
        ws[f"D{r}"].number_format = '#,##0.00'

for sheet in ("Sample", "Template"):
    del wb[sheet]
wb.save(out)
print("written", out, [w.title for w in wb.worksheets])
