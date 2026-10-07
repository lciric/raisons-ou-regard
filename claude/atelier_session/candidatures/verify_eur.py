import sys
from openpyxl import load_workbook
p = sys.argv[1]
v, f = load_workbook(p, data_only=True), load_workbook(p)
for name in ("Minimum", "Mainline", "Maximum"):
    s, fs = v[name], f[name]
    rate = s["B1"].value
    eur_lines = []
    print(f"== {name}: B1={rate:.6f} EUR/USD ({fs['B1'].value}); contingency {s['B2'].value}")
    tot_eur = 0.0
    for r in range(6, 20):
        if s[f"A{r}"].value is None:
            continue
        d = s[f"D{r}"].value or 0
        tot_eur += d
        print(f"   r{r} {s[f'A{r}'].value:10} D={d:10.2f} EUR  F={s[f'F{r}'].value:10.2f} USD  E={s[f'E{r}'].value:6.1%}  formula={fs[f'D{r}'].value}")
    sub_usd, cont_usd, tot_usd = s["F3"].value, s["F2"].value, s["F1"].value
    print(f"   subtotal {tot_eur:,.2f} EUR = {sub_usd:,.2f} USD; contingency {cont_usd:,.2f} USD; TOTAL {tot_usd:,.2f} USD = {tot_usd*rate:,.2f} EUR")
