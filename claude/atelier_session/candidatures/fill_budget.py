"""Fills the EA Funds budget template with the three scenarios of the application (minimum, mainline, maximum).
Writes only the template's input cells (columns A to D of rows 6 on, and the contingency B2); keeps its formulas."""
import sys
from copy import copy
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

src, out = sys.argv[1], sys.argv[2]
wb = load_workbook(src)
tpl = wb["Template"]

STIPEND = ("Personnel", "Stipend for the applicant: [[months]] x [[USD per month]]",
           "[[To fill by the applicant: gross amount, including income tax and social charges; leave empty if no stipend]]", None)
API_DATA = ("Other", "Claude API: generation of the training data of the six arms",
            "Programme estimate: 500-1,400 USD at standard prices; upper end (less with batch pricing)", 1400)
RESERVE = ("Other", "GPU reserve: 720 GPU-hours of H100 SXM at 3.50 USD/hour, for the larger gaze test",
           "The gaze test grows from 150-280 to 600-1,000 GPU-h at 400 scenarios x 10 generations; spent only if the pilot's power simulation calls for it",
           "=720*3.5")
SCENARIOS = {
    "Minimum": [
        STIPEND,
        ("Other", "GPU compute: 560 GPU-hours of H100 SXM rental (vast.ai) at 3.50 USD/hour",
         "Upper end of the programme's estimate for the first paper: minimal decisive experiment and localization (270-560 GPU-h). H100 SXM prices seen on 6 Oct 2026: 2.86-4.04 USD/h",
         "=560*3.5"),
        API_DATA,
        ("Other", "Claude API: cue sets, held-out scenarios, judges on stratified subsamples",
         "Estimate for the first paper; about 1,000 USD less if Anthropic's External Researcher Access Program grants API credits (applied for separately)", 1000),
    ],
    "Mainline": [
        STIPEND,
        ("Other", "GPU compute: 560 GPU-hours of H100 SXM rental (vast.ai) at 3.50 USD/hour",
         "Upper end of the programme's estimate for the first paper: minimal decisive experiment and localization (270-560 GPU-h). H100 SXM prices seen on 6 Oct 2026: 2.86-4.04 USD/h",
         "=560*3.5"),
        RESERVE,
        API_DATA,
        ("Other", "Claude API: cue sets, held-out scenarios, judges on stratified subsamples",
         "Estimate for the first paper; about 1,000 USD less if Anthropic's External Researcher Access Program grants API credits (applied for separately)", 1000),
    ],
    "Maximum": [
        STIPEND,
        ("Other", "GPU compute: 1,960 GPU-hours of H100 SXM rental (vast.ai) at 3.50 USD/hour",
         "Upper end of the programme's estimate for all its phases, without the deferred distress phase and the optional analyses (980-1,960 GPU-h). H100 SXM prices seen on 6 Oct 2026: 2.86-4.04 USD/h",
         "=1960*3.5"),
        RESERVE,
        API_DATA,
        ("Other", "Claude API: cue sets, held-out scenarios, judges and ground truth for all phases",
         "Programme estimate: up to about 3,000 USD; about 1,000 USD less if Anthropic's External Researcher Access Program grants API credits (applied for separately)", 3000),
    ],
}

for name, lines in SCENARIOS.items():
    ws = wb.copy_worksheet(tpl)
    ws.title = name
    for col, dim in tpl.column_dimensions.items():
        ws.column_dimensions[col].width = dim.width
    dv = DataValidation(type="list", formula1='"Personnel,Materials,Operations,Travel,Other"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A6:A41")
    ws["B2"] = 0.10      # contingency: the application form suggests about 10 %, the template 5 %
    for i, (kind, detail, note, amount) in enumerate(lines):
        r = 6 + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = kind, detail, note
        if amount is not None:
            ws[f"D{r}"] = amount
            ws[f"D{r}"].number_format = '#,##0.00'

for sheet in ("Sample", "Template"):
    del wb[sheet]
wb.save(out)
print("written", out, [ws.title for ws in wb.worksheets])
