import sys
from openpyxl import load_workbook
p = sys.argv[1]
vals, forms = load_workbook(p, data_only=True), load_workbook(p)
for name in ("Minimum", "Mainline", "Maximum"):
    v, f = vals[name], forms[name]
    print(f"== {name}: subtotal F3={v['F3'].value}  contingency F2={v['F2'].value} (B2={v['B2'].value})  total F1={v['F1'].value}")
    print("   formulas:", f["F1"].value, f["F2"].value, f["F3"].value, f["F7"].value, f["E7"].value)
    for r in range(6, 12):
        if v[f"A{r}"].value or v[f"D{r}"].value:
            print(f"   r{r} {v[f'A{r}'].value:9} D={v[f'D{r}'].value} F={v[f'F{r}'].value} E={round(v[f'E{r}'].value or 0, 3)} | {str(v[f'B{r}'].value)[:70]}")
    print("   validations:", [(d.formula1, str(d.sqref)) for d in f.data_validations.dataValidation])
